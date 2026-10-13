"""Tests for document text extraction, cleaning, chunking, and the ingestion pipeline."""

import io
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch
import docx
import pytest

from app.core.errors import AppException, BadRequestException
from app.models.knowledge import DocumentSourceType, DocumentStatus, IngestionResult
from app.services.embedding_service import EmbeddingService, MockEmbeddingProvider
from app.services.ingestion_service import IngestionService
from app.services.vector_search_service import VectorSearchService
from app.utils.text_cleaner import clean_text
from app.utils.text_chunker import chunk_text, count_tokens
from app.utils.text_extractor import (
    extract_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_from_txt,
)


# ==============================================================================
# 1. Text Cleaning Tests
# ==============================================================================


def test_clean_text_normalizes_whitespace_and_newlines() -> None:
    """Verify clean_text normalizes irregular line endings, tabs, and excess blank lines."""
    raw = "  Line 1  with   spaces.  \r\n\r\n\r\n\r\nLine   2 with \t tabs.\r\n"
    cleaned = clean_text(raw)
    assert "Line 1 with spaces." in cleaned
    assert "Line 2 with tabs." in cleaned
    # Should not have more than 2 consecutive newlines
    assert "\n\n\n" not in cleaned


def test_clean_text_strips_control_characters() -> None:
    """Verify clean_text removes non-printable control characters while keeping standard chars."""
    raw = "Temple\x00 Information\x08 and\x0c History."
    cleaned = clean_text(raw)
    assert cleaned == "Temple Information and History."


def test_clean_text_handles_empty() -> None:
    """Verify clean_text handles empty or whitespace-only strings."""
    assert clean_text("") == ""
    assert clean_text("   \n\t  ") == ""


# ==============================================================================
# 2. Text Chunking Tests
# ==============================================================================


def test_chunk_text_single_short_string() -> None:
    """Verify short text within chunk_size is returned as a single chunk."""
    text = "Sri Venkateswara Temple is an ancient Vaishnavite temple."
    chunks = chunk_text(text, chunk_size=200, chunk_overlap=20)
    assert len(chunks) == 1
    assert chunks[0] == text


def test_chunk_text_splits_on_boundaries() -> None:
    """Verify chunk_text splits along natural sentence or paragraph boundaries."""
    paragraph1 = "First paragraph details the morning pooja rituals starting at 5:00 AM."
    paragraph2 = "Second paragraph discusses the evening darshan queue timings and ticket prices."
    text = f"{paragraph1}\n\n{paragraph2}"

    chunks = chunk_text(text, chunk_size=80, chunk_overlap=15)
    assert len(chunks) >= 2
    assert "First paragraph" in chunks[0]
    assert "Second paragraph" in chunks[1]


def test_chunk_text_validation_errors() -> None:
    """Verify chunk_text raises BadRequestException on invalid size or overlap."""
    with pytest.raises(BadRequestException):
        chunk_text("Some text", chunk_size=0)

    with pytest.raises(BadRequestException):
        chunk_text("Some text", chunk_size=100, chunk_overlap=-5)

    with pytest.raises(BadRequestException):
        chunk_text("Some text", chunk_size=100, chunk_overlap=100)


def test_count_tokens_calculation() -> None:
    """Verify count_tokens returns positive estimate."""
    assert count_tokens("") == 0
    assert count_tokens("Temple darshan queue") >= 3


# ==============================================================================
# 3. Text Extraction Tests
# ==============================================================================


def test_extract_text_txt() -> None:
    """Verify extract_text from plain text strings and bytes."""
    text = "Darshan timings: 06:00 to 21:00."
    assert extract_text_from_txt(text) == text
    assert extract_text_from_txt(text.encode("utf-8")) == text
    assert extract_text(text, file_type="txt") == text


def test_extract_text_docx() -> None:
    """Verify extract_text_from_docx extracts paragraphs and tables from real DOCX bytes."""
    doc = docx.Document()
    doc.add_heading("Temple Architecture", 0)
    doc.add_paragraph("The gopuram stands 50 meters tall.")
    table = doc.add_table(rows=1, cols=2)
    row = table.rows[0]
    row.cells[0].text = "Morning Pooja"
    row.cells[1].text = "06:00 AM"

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    extracted = extract_text_from_docx(docx_bytes)
    assert "Temple Architecture" in extracted
    assert "The gopuram stands 50 meters tall." in extracted
    assert "Morning Pooja | 06:00 AM" in extracted


def test_extract_text_pdf_mock() -> None:
    """Verify extract_text_from_pdf calls PdfReader and extracts page text."""
    with patch("app.utils.text_extractor.PdfReader") as mock_pdf_reader:
        mock_instance = MagicMock()
        mock_instance.is_encrypted = False
        mock_page1 = MagicMock()
        mock_page1.extract_text.return_value = "Page 1: Temple Inscriptions."
        mock_page2 = MagicMock()
        mock_page2.extract_text.return_value = "Page 2: Chola Dynasty Contributions."
        mock_instance.pages = [mock_page1, mock_page2]
        mock_pdf_reader.return_value = mock_instance

        res = extract_text_from_pdf(b"%PDF-1.4 mock content")
        assert "Page 1: Temple Inscriptions." in res
        assert "Page 2: Chola Dynasty Contributions." in res


def test_extract_text_unsupported_format() -> None:
    """Verify extract_text rejects unsupported document formats."""
    with pytest.raises(BadRequestException):
        extract_text(b"some content", file_type="xlsx")


# ==============================================================================
# 4. Ingestion Service Pipeline Tests
# ==============================================================================


@pytest.fixture
def test_ingestion_service() -> IngestionService:
    """Fixture providing IngestionService with MockEmbeddingProvider."""
    embed_svc = EmbeddingService(MockEmbeddingProvider())
    vector_svc = VectorSearchService(embed_service=embed_svc)
    return IngestionService(embed_service=embed_svc, vector_service=vector_svc)


@pytest.mark.anyio
async def test_ingest_document_validation(test_ingestion_service: IngestionService) -> None:
    """Verify validation on empty templeId, title, or unreadable document."""
    with pytest.raises(BadRequestException):
        await test_ingestion_service.ingest_document(temple_id="", source="content", title="Title")

    with pytest.raises(BadRequestException):
        await test_ingestion_service.ingest_document(temple_id="t1", source="content", title="")

    with pytest.raises(BadRequestException):
        await test_ingestion_service.ingest_document(temple_id="t1", source="   ", title="Title")


@pytest.mark.anyio
async def test_ingest_document_success(test_ingestion_service: IngestionService) -> None:
    """Verify full successful ingestion flow with document creation and vector chunks."""
    temple_id = "temple_madurai"
    title = "Meenakshi Temple History"
    text_content = (
        "Meenakshi Amman Temple is a historic Hindu temple located on the southern bank of the Vaigai River. "
        "It is dedicated to Meenakshi, a form of Parvati, and her consort, Sundareswarar. "
        "The temple was reconstructed in the 16th and 17th centuries by the Nayak dynasty."
    )

    mock_doc_col = MagicMock()
    mock_doc_col.find_one = AsyncMock(return_value=None)
    mock_doc_col.insert_one = AsyncMock(return_value=MagicMock(inserted_id="670ab349f123456789abcdef"))
    mock_doc_col.update_one = AsyncMock(return_value=MagicMock(modified_count=1))

    mock_chunks_col = MagicMock()
    mock_chunks_col.insert_many = AsyncMock(
        return_value=MagicMock(inserted_ids=["chunk_1", "chunk_2"])
    )

    with patch.object(IngestionService, "documents_collection", new=mock_doc_col), patch.object(
        VectorSearchService, "chunks_collection", new=mock_chunks_col
    ):
        result = await test_ingestion_service.ingest_document(
            temple_id=temple_id,
            source=text_content,
            title=title,
            source_type=DocumentSourceType.TEXT,
            category="history",
            chunk_size=120,
            chunk_overlap=20,
            metadata={"sourceFile": "meenakshi.txt"},
        )

        assert isinstance(result, IngestionResult)
        assert result.temple_id == temple_id
        assert result.title == title
        assert result.status == DocumentStatus.PROCESSED
        assert result.chunks_created >= 1
        assert result.is_duplicate is False
        assert result.content_hash is not None

        # Verify document was inserted with status PROCESSING and updated to PROCESSED
        mock_doc_col.insert_one.assert_called_once()
        mock_doc_col.update_one.assert_called_once()
        mock_chunks_col.insert_many.assert_called_once()

        stored_chunks = mock_chunks_col.insert_many.call_args[0][0]
        assert len(stored_chunks) == result.chunks_created
        for c in stored_chunks:
            assert c["templeId"] == temple_id
            assert c["documentId"] == "670ab349f123456789abcdef"
            assert c["title"] == title
            assert c["category"] == "history"
            assert "embedding" in c
            assert len(c["embedding"]) == 384


@pytest.mark.anyio
async def test_ingest_document_deduplication(test_ingestion_service: IngestionService) -> None:
    """Verify deduplication detects identical content and skips redundant ingestion."""
    temple_id = "temple_somnath"
    title = "Somnath Aarti Timings"
    text = "Morning Aarti is held at 7:00 AM. Evening Aarti is at 7:00 PM."

    existing_doc_record = {
        "_id": "existing_doc_id_999",
        "templeId": temple_id,
        "title": title,
        "status": DocumentStatus.PROCESSED.value,
        "chunkCount": 2,
    }

    mock_doc_col = MagicMock()
    # Return existing document when checking contentHash
    mock_doc_col.find_one = AsyncMock(return_value=existing_doc_record)

    with patch.object(IngestionService, "documents_collection", new=mock_doc_col):
        # 1. Without force_reingest: should skip
        res = await test_ingestion_service.ingest_document(
            temple_id=temple_id,
            source=text,
            title=title,
            force_reingest=False,
        )
        assert res.is_duplicate is True
        assert res.document_id == "existing_doc_id_999"
        assert res.chunks_created == 2


@pytest.mark.anyio
async def test_ingest_document_failure_rollback(test_ingestion_service: IngestionService) -> None:
    """Verify clean failure handling updates document to FAILED and cleans up partial chunks."""
    temple_id = "temple_kashi"
    title = "Kashi Vishwanath Corridor"
    text = "Information regarding Kashi Vishwanath temple corridor and darshan queues."

    mock_doc_col = MagicMock()
    mock_doc_col.find_one = AsyncMock(return_value=None)
    mock_doc_col.insert_one = AsyncMock(return_value=MagicMock(inserted_id="doc_fail_123"))
    mock_doc_col.update_one = AsyncMock(return_value=MagicMock())

    # Simulate vector storage error
    mock_vector_svc = MagicMock()
    mock_vector_svc.store_chunks = AsyncMock(side_effect=RuntimeError("Vector storage failure"))
    mock_vector_svc.delete_chunks_by_document = AsyncMock(return_value=0)

    test_ingestion_service._vector_service = mock_vector_svc

    with patch.object(IngestionService, "documents_collection", new=mock_doc_col):
        with pytest.raises(AppException) as exc_info:
            await test_ingestion_service.ingest_document(
                temple_id=temple_id,
                source=text,
                title=title,
            )

        assert "Document ingestion failed" in str(exc_info.value)
        # Verify rollback attempted
        mock_vector_svc.delete_chunks_by_document.assert_called_once_with(temple_id, "doc_fail_123")
        # Verify document updated with FAILED status
        mock_doc_col.update_one.assert_called_once()
        update_args = mock_doc_col.update_one.call_args[0][1]["$set"]
        assert update_args["status"] == DocumentStatus.FAILED.value


@pytest.mark.anyio
async def test_ingest_convenience_methods(test_ingestion_service: IngestionService, tmp_path: Path) -> None:
    """Verify ingest_text and ingest_file convenience methods."""
    sample_file = tmp_path / "temple_guide.txt"
    sample_file.write_text("Guide to visiting the sanctum sanctorum.", encoding="utf-8")

    mock_res = IngestionResult(
        document_id="doc_conv_1",
        temple_id="temple_1",
        title="Temple Guide",
        status=DocumentStatus.PROCESSED,
        chunks_created=1,
    )

    with patch.object(IngestionService, "ingest_document", new=AsyncMock(return_value=mock_res)) as mock_ingest:
        res1 = await test_ingestion_service.ingest_file(
            file_path=sample_file,
            temple_id="temple_1",
            category="general",
        )
        assert res1.document_id == "doc_conv_1"
        mock_ingest.assert_called_once()

        res2 = await test_ingestion_service.ingest_text(
            text="FAQ: Is photography allowed? No.",
            temple_id="temple_1",
            title="Photography FAQ",
            category="faq",
        )
        assert res2.document_id == "doc_conv_1"
