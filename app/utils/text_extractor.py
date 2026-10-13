"""Text extraction utilities supporting PDF, DOCX, and TXT documents."""

import io
from pathlib import Path
from typing import Optional, Union
import docx
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from app.core.errors import BadRequestException
from app.core.logging import get_logger

logger = get_logger(__name__)


def extract_text_from_pdf(source: Union[bytes, io.BytesIO, Path, str]) -> str:
    """Extract plain text from a PDF document.

    Args:
        source: Raw PDF bytes, file path, or BytesIO buffer.

    Returns:
        str: Extracted text content across all pages.
    """
    try:
        stream = io.BytesIO(source) if isinstance(source, bytes) else source
        reader = PdfReader(stream)

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise BadRequestException("PDF document is encrypted and cannot be processed.")

        pages_text = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages_text.append(page_text.strip())

        extracted = "\n\n".join(pages_text)
        if not extracted.strip():
            logger.warning("PDF extraction yielded empty text (may be image-based or scanned).")
        return extracted
    except PyPdfError as exc:
        raise BadRequestException(f"Failed to read PDF file: {str(exc)}")
    except Exception as exc:
        if isinstance(exc, BadRequestException):
            raise
        raise BadRequestException(f"Unexpected error while extracting PDF: {str(exc)}")


def extract_text_from_docx(source: Union[bytes, io.BytesIO, Path, str]) -> str:
    """Extract plain text from a Word (.docx) document including tables.

    Args:
        source: Raw DOCX bytes, file path, or BytesIO buffer.

    Returns:
        str: Extracted text from paragraphs and tables.
    """
    try:
        stream = io.BytesIO(source) if isinstance(source, bytes) else source
        doc = docx.Document(stream)

        content_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_parts.append(text)

        # Extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    content_parts.append(" | ".join(row_cells))

        return "\n\n".join(content_parts)
    except Exception as exc:
        raise BadRequestException(f"Failed to extract DOCX document: {str(exc)}")


def extract_text_from_txt(source: Union[bytes, Path, str]) -> str:
    """Extract plain text from a text file or raw bytes with encoding fallbacks.

    Args:
        source: Raw bytes, file path, or string.

    Returns:
        str: Decoded text content.
    """
    if isinstance(source, str):
        path = Path(source)
        if path.exists() and path.is_file():
            source = path.read_bytes()
        else:
            return source

    if isinstance(source, Path):
        source = source.read_bytes()

    if isinstance(source, bytes):
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252", "utf-16"):
            try:
                return source.decode(encoding)
            except (UnicodeDecodeError, LookupError):
                continue
        raise BadRequestException("Unable to decode text document using standard encodings.")

    return str(source)


def extract_text(
    source: Union[bytes, Path, str],
    file_type: Optional[str] = None,
    filename: Optional[str] = None,
) -> str:
    """Unified extractor detecting document type from file extension or explicit type parameter.

    Args:
        source: Raw bytes, Path, or file content.
        file_type: Optional explicit type ('pdf', 'docx', 'txt', 'text', 'md').
        filename: Optional filename hint to deduce file extension.

    Returns:
        str: Extracted text content.
    """
    if hasattr(file_type, "value"):
        resolved_type = str(file_type.value).lower().strip()
    else:
        resolved_type = str(file_type or "").lower().strip()

    if "." in resolved_type and not filename:
        resolved_type = resolved_type.split(".")[-1]

    if not resolved_type:
        if filename:
            ext = Path(filename).suffix.lower().lstrip(".")
            resolved_type = ext
        elif isinstance(source, (Path, str)):
            ext = Path(source).suffix.lower().lstrip(".")
            if ext:
                resolved_type = ext

    # Fallback to txt if unresolved
    if not resolved_type:
        resolved_type = "txt"

    if resolved_type in ("pdf",):
        return extract_text_from_pdf(source)
    elif resolved_type in ("docx", "doc"):
        return extract_text_from_docx(source)
    elif resolved_type in ("txt", "text", "md", "markdown", "faq"):
        return extract_text_from_txt(source)
    else:
        raise BadRequestException(
            f"Unsupported document format '{resolved_type}'. Supported formats: PDF, DOCX, TXT."
        )
